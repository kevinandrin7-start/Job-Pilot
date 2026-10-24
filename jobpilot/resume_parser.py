"""Resume text extraction (PDF, DOCX, TXT) and LLM parsing into structured JSON."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from .llm import LLMBackend

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ResumeError(ValueError):
    pass


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif ext == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif ext == ".txt":
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
    else:
        raise ResumeError(f"Unsupported file type {ext!r}. Upload a PDF, DOCX or TXT file.")
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ResumeError("No text could be extracted. If this is a scanned PDF, upload a DOCX or TXT instead.")
    return text


_S = {"type": "string"}
_SA = {"type": "array", "items": _S}


def _obj(props: dict) -> dict:
    return {"type": "object", "properties": props, "required": list(props), "additionalProperties": False}


RESUME_SCHEMA = _obj({
    "name": _S,
    "contact": _obj({"email": _S, "phone": _S, "location": _S, "links": _SA}),
    "summary": _S,
    "skills": _SA,
    "experience": {"type": "array", "items": _obj({
        "title": _S, "company": _S, "location": _S, "start_date": _S, "end_date": _S, "bullets": _SA,
    })},
    "education": {"type": "array", "items": _obj({
        "institution": _S, "degree": _S, "field": _S, "start_date": _S, "end_date": _S, "details": _SA,
    })},
    "projects": {"type": "array", "items": _obj({
        "name": _S, "description": _S, "technologies": _SA, "bullets": _SA,
    })},
    "certifications": {"type": "array", "items": _obj({"name": _S, "issuer": _S, "date": _S})},
})

PARSE_SYSTEM = """You convert resume text into structured JSON.

Rules:
- Copy facts exactly as written: names, employers, job titles, dates, numbers, degrees, skills and bullet wording.
- Do not summarize, embellish, correct, infer or add anything that is not in the text.
- If a field is missing in the resume, use an empty string or an empty list.
- Keep every bullet point, in its original order, as its own list item.
- "summary" is the resume's own summary/objective paragraph, copied verbatim, or "" if there is none.
- "skills" is the list of skills exactly as the resume lists them, one item per skill."""


def empty_resume() -> dict:
    return {
        "name": "", "contact": {"email": "", "phone": "", "location": "", "links": []},
        "summary": "", "skills": [], "experience": [], "education": [], "projects": [], "certifications": [],
    }


def normalize_resume(data: dict) -> dict:
    """Fill in any missing keys so templates and the tailor can rely on the shape."""
    out = empty_resume()
    if not isinstance(data, dict):
        return out
    out["name"] = str(data.get("name") or "")
    contact = data.get("contact") or {}
    for k in ("email", "phone", "location"):
        out["contact"][k] = str(contact.get(k) or "")
    out["contact"]["links"] = [str(x) for x in contact.get("links") or [] if x]
    out["summary"] = str(data.get("summary") or "")
    out["skills"] = [str(s).strip() for s in data.get("skills") or [] if str(s).strip()]
    for e in data.get("experience") or []:
        out["experience"].append({
            "title": str(e.get("title") or ""), "company": str(e.get("company") or ""),
            "location": str(e.get("location") or ""), "start_date": str(e.get("start_date") or ""),
            "end_date": str(e.get("end_date") or ""),
            "bullets": [str(b) for b in e.get("bullets") or [] if str(b).strip()],
        })
    for e in data.get("education") or []:
        out["education"].append({
            "institution": str(e.get("institution") or ""), "degree": str(e.get("degree") or ""),
            "field": str(e.get("field") or ""), "start_date": str(e.get("start_date") or ""),
            "end_date": str(e.get("end_date") or ""),
            "details": [str(b) for b in e.get("details") or [] if str(b).strip()],
        })
    for p in data.get("projects") or []:
        out["projects"].append({
            "name": str(p.get("name") or ""), "description": str(p.get("description") or ""),
            "technologies": [str(t) for t in p.get("technologies") or [] if str(t).strip()],
            "bullets": [str(b) for b in p.get("bullets") or [] if str(b).strip()],
        })
    for c in data.get("certifications") or []:
        out["certifications"].append({
            "name": str(c.get("name") or ""), "issuer": str(c.get("issuer") or ""), "date": str(c.get("date") or ""),
        })
    return out


def parse_resume(llm: LLMBackend, resume_text: str) -> dict:
    user = f"Resume text:\n<resume>\n{resume_text}\n</resume>\n\nReturn the structured JSON."
    data = llm.json(PARSE_SYSTEM, user, RESUME_SCHEMA)
    return normalize_resume(data)


def resume_to_text(resume: dict) -> str:
    """Flatten structured resume JSON to plain text (used for matching and fact checks)."""
    return json.dumps(resume, ensure_ascii=False)
