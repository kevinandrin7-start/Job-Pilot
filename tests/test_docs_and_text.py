"""DOCX rendering, resume text extraction, HTML-to-text and source parsing."""

from __future__ import annotations

from docx import Document

from conftest import PARSED_RESUME, REMOTIVE_PAYLOAD, remotive_with_mock, write_minimal_pdf

from jobpilot.job_sources import SearchQuery, html_to_text, parse_salary
from jobpilot.matcher import passes_filters
from jobpilot.resume_parser import extract_text, normalize_resume
from jobpilot.tailor import render_cover_docx, render_resume_docx


def _docx_text(path) -> str:
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


def test_render_resume_docx(tmp_path):
    path = render_resume_docx(PARSED_RESUME, tmp_path / "resume.docx")
    text = _docx_text(path)
    for expected in ("Jordan Rivera", "jordan@example.com", "EXPERIENCE", "Data Analyst – Acme Corp",
                     "Built 12 Tableau dashboards used by 300 sales staff", "B.S., Statistics – University of Texas",
                     "Tableau Desktop Specialist, Tableau, 2022"):
        assert expected in text
    bullets = [p for p in Document(str(path)).paragraphs if p.style.name == "List Bullet"]
    assert len(bullets) >= 3


def test_render_cover_docx(tmp_path):
    path = render_cover_docx("Dear Hiring Team,\n\nPara one.\n\nPara two.", PARSED_RESUME, tmp_path / "c.docx")
    text = _docx_text(path)
    assert "Jordan Rivera" in text and "Para two." in text


def test_extract_text_formats(tmp_path):
    txt = tmp_path / "r.txt"
    txt.write_text("Jordan Rivera\nSQL")
    assert "SQL" in extract_text(txt)

    docx_path = render_resume_docx(PARSED_RESUME, tmp_path / "r.docx")
    assert "Acme Corp" in extract_text(docx_path)

    pdf = tmp_path / "r.pdf"
    write_minimal_pdf(pdf, ["Jordan Rivera", "Data Analyst at Acme Corp"])
    assert "Acme Corp" in extract_text(pdf)


def test_html_to_text():
    html = ("<html><head><style>p{}</style><script>alert(1)</script></head><body><h2>About</h2>"
            "<p>We build&nbsp;tools &amp; things.</p><ul><li>SQL</li><li>Python</li></ul></body></html>")
    text = html_to_text(html)
    assert "alert" not in text and "p{}" not in text
    assert "We build tools & things." in text
    assert "- SQL" in text and "- Python" in text
    assert text.startswith("About")


def test_parse_salary():
    assert parse_salary("$90k - $120k") == (90000, 120000)
    assert parse_salary("100,000 USD") == (100000, 100000)
    assert parse_salary("competitive") == (None, None)


def test_remotive_parsing(settings):
    jobs = remotive_with_mock(settings).search(SearchQuery(keywords="data analyst"))
    assert [j.external_id for j in jobs] == ["101", "102"]
    assert jobs[0].company == "Initrode" and jobs[0].remote is True
    assert jobs[0].salary_min == 100000 and "Dashboards" in jobs[0].description


def test_filters():
    prefs = {"excluded_keywords": ["clearance"], "remote_only": True, "min_salary": 90000}
    ok, _ = passes_filters({"title": "Analyst", "remote": 1, "description": "SQL", "salary_max": 100000}, prefs)
    assert ok
    assert not passes_filters({"title": "Analyst", "remote": 1, "description": "needs Clearance"}, prefs)[0]
    assert not passes_filters({"title": "Analyst", "remote": 0, "location": "Austin"}, prefs)[0]
    assert not passes_filters({"title": "Analyst", "remote": 1, "salary_max": 50000}, prefs)[0]


def test_normalize_fills_missing():
    r = normalize_resume({"name": "X", "experience": [{"title": "T"}]})
    assert r["experience"][0]["bullets"] == [] and r["contact"]["links"] == []
